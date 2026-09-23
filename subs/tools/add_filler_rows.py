# -*- coding: utf-8 -*-
"""Insert ALL no-sheet lines (sung ad-lib + spoken chatter) into the two Frau
spreadsheets, in a new 「插白补充」 sheet; original sheets untouched.

- Dance with me  : sung ad-lib (user OK'd sweep) -> Chinese 和我一起跳吧
- 6 chatter lines: spoken -> Chinese, static
- Tu Tu Ru vocalise: keep CoZ romaji verbatim (listed for completeness)
"""
import os
import openpyxl
from openpyxl.styles import Font, PatternFill

ROOT = r"D:\DATA\tran\agent tran\9.6文本外工作"

VOC = "Tu Tu Ru Tu Ru Tu♪ Tu Tu Ru Tu Ru Tu♪"
VOC_B = "Tu Tu Ru Tu Ru Tu♪ Da-li-da-li-ya-a-a♪"
VOC_C = "Tu Tu Ru La La La La♪"

JOBS = [
    ("歌词翻译/edfrau_歌词对照.xlsx", [
        # --- sung ad-lib ---
        ("0:00:45.04", "0:00:46.49", "Dance with me", "Dance with me", "和我一起跳吧",
         "英文和声（唱）。CoZ自加，日文层无对应行。带扫色。", "和我一起跳吧"),
        ("0:00:59.83", "0:01:01.26", "Dance with me", "Dance with me", "和我一起跳吧",
         "同上，第二遍。", "和我一起跳吧"),
        # --- vocalise: keep original romaji ---
        ("0:00:02.14", "0:00:05.94", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "拟声垫曲。CoZ的translation层即放罗马音；补丁逐字保留原版，静态。", VOC),
        ("0:00:05.80", "0:00:09.61", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:00:09.52", "0:00:13.28", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:00:13.23", "0:00:16.19", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
        ("0:01:19.68", "0:01:23.42", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:01:23.39", "0:01:27.13", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:01:27.06", "0:01:30.81", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:01:30.78", "0:01:33.88", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
    ]),
    ("歌词翻译/livedance_歌词对照.xlsx", [
        # --- sung ad-lib ---
        ("0:00:45.44", "0:00:46.89", "Dance with me", "Dance with me", "和我一起跳吧",
         "英文和声（唱）。CoZ自加。带扫色。", "和我一起跳吧"),
        ("0:01:00.23", "0:01:01.66", "Dance with me", "Dance with me", "和我一起跳吧",
         "同上，第二遍。", "和我一起跳吧"),
        ("0:02:14.14", "0:02:15.56", "Dance with me", "Dance with me", "和我一起跳吧",
         "同上，第三遍。", "和我一起跳吧"),
        # --- spoken chatter (static) ---
        ("0:01:13.03", "0:01:15.10", "（无日文原文，现场口白）", "Everybody now!", "大家一起来！",
         "CoZ听的现场喊话。说，不是唱；静态不扫色。", "大家一起来！"),
        ("0:01:24.14", "0:01:27.51", "（无日文原文，现场口白）", "Oh man, now I really wanna dance too...",
         "糟了，我也真想跳起来……", "同上。", "糟了，我也真想跳起来……"),
        ("0:01:27.31", "0:01:28.66", "（无日文原文，现场口白）", "Frau-tan, c'mon!", "芙劳炭，来吧！",
         "同上。称呼按术语表『芙劳炭』（隔壁作『芙兰碳』）。", "芙劳炭，来吧！"),
        ("0:02:28.08", "0:02:30.27", "（无日文原文，现场口白）", "Enako-chan, come dance with us!",
         "Enako酱来一起跳吧", "同上。『Enako』为CoZ听写的呼喊（游戏文本查无此名，但音频确有喊名字）；用户拍板保留并定稿措辞。", "Enako酱来一起跳吧"),
        ("0:02:43.67", "0:02:45.10", "（无日文原文，现场口白）", "One more time!", "再来一次！",
         "同上。", "再来一次！"),
        ("0:03:01.16", "0:03:05.22", "（无日文原文，现场口白）",
         "Ah, I'm beat... this burnt way more\\N..calories than I thought",
         "啊，累死了……消耗的热量\\N比想象中多太多了。", "同上。\\N为分行。", "啊，累死了……消耗的热量\\N比想象中多太多了。"),
        # --- vocalise ---
        ("0:00:02.54", "0:00:06.34", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "拟声垫曲，保留原版罗马音，静态。", VOC),
        ("0:00:06.20", "0:00:10.01", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:00:09.92", "0:00:13.68", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:00:13.63", "0:00:16.59", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
        ("0:01:16.43", "0:01:20.17", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:01:20.14", "0:01:23.88", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:01:23.81", "0:01:27.56", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:01:27.53", "0:01:30.63", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
        ("0:01:50.28", "0:01:54.02", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:01:53.95", "0:01:57.71", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:01:57.65", "0:02:01.44", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:02:01.35", "0:02:04.30", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
        ("0:02:30.28", "0:02:34.02", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:02:33.95", "0:02:37.71", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:02:37.65", "0:02:41.44", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:02:41.35", "0:02:44.30", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
        ("0:02:45.03", "0:02:48.82", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:02:48.77", "0:02:52.53", "（无日文原文，为旋律拟声）", VOC_B, "（不译，保留原版罗马音）",
         "同上。", VOC_B),
        ("0:02:52.47", "0:02:56.17", "（无日文原文，为旋律拟声）", VOC, "（不译，保留原版罗马音）",
         "同上。", VOC),
        ("0:02:56.11", "0:02:59.23", "（无日文原文，为旋律拟声）", VOC_C, "（不译，保留原版罗马音）",
         "同上。", VOC_C),
    ]),
]

HDR = ("起始时间", "结束时间", "日文原文", "罗马音", "中文译文", "翻译说明（逐词+斟酌）", "补丁内实际显示")
HDR_FILL = PatternFill("solid", fgColor="FFF2CC")
CAT_SUNG = PatternFill("solid", fgColor="D9EAD3")   # green: sung, swept
CAT_TALK = PatternFill("solid", fgColor="FCE5CD")   # orange: spoken, static
CAT_VOC = PatternFill("solid", fgColor="D0E0E3")    # teal: keep romaji
HDR_FONT = Font(bold=True)

for rel, rows in JOBS:
    src = os.path.join(ROOT, rel)
    dst = os.path.join(ROOT, rel.replace(".xlsx", "_插白补充.xlsx"))
    wb = openpyxl.load_workbook(src)
    if "插白补充" in wb.sheetnames:
        del wb["插白补充"]
    ws = wb.create_sheet("插白补充")
    ws.append(HDR)
    for c in ws[1]:
        c.font = HDR_FONT
        c.fill = HDR_FILL
    for st, en, jp, rm, zh, note, shown in rows:
        ws.append((st, en, jp, rm, zh, note, shown))
        ridx = ws.max_row
        if "Dance with me" in rm:
            fill = CAT_SUNG
        elif "Tu Tu Ru" in rm:
            fill = CAT_VOC
        else:
            fill = CAT_TALK
        for c in ws[ridx]:
            c.fill = fill
    for col, w in zip("ABCDEFG", (11, 11, 26, 40, 24, 56, 24)):
        ws.column_dimensions[col].width = w
    wb.save(dst)
    print("wrote", os.path.relpath(dst, ROOT), "| rows:", ws.max_row - 1)
